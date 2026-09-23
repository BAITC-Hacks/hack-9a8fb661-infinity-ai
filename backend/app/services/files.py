"""Вложения к диалогу: чтение документов сотрудников (docx, pdf, xlsx, csv, изображения).

Файлы держатся в памяти процесса на сессию диалога (не пишутся на диск), текст обрезается
до MAX_CHARS, изображения передаются в LLM с поддержкой изображений (OpenAI) как data URL.
"""
import base64
import io
import time
import uuid

import pandas as pd

MAX_BYTES = 15 * 1024 * 1024
MAX_CHARS = 20000
ALLOWED = {".docx", ".doc", ".pdf", ".xlsx", ".xls", ".csv", ".jpg", ".jpeg", ".png", ".txt"}
_FILES: dict[str, dict] = {}
TTL_S = 6 * 3600


class FileError(ValueError):
    pass


def _table_text(df: pd.DataFrame, name: str) -> str:
    df = df.dropna(how="all").dropna(axis=1, how="all")
    parts = [f"Таблица «{name}»: {len(df)} строк × {df.shape[1]} столбцов. Столбцы: {', '.join(map(str, df.columns))}."]
    num = df.select_dtypes("number")
    if not num.empty:
        parts.append("Сводка по числовым столбцам:\n" + num.describe().T[["mean", "min", "max"]].round(3).to_string())
    parts.append("Первые строки:\n" + df.head(25).to_string(max_colwidth=40))
    return "\n".join(parts)


def extract(name: str, data: bytes) -> dict:
    ext = "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in ALLOWED:
        raise FileError(f"Формат {ext or '?'} не поддерживается. Можно: {', '.join(sorted(ALLOWED))}")
    if len(data) > MAX_BYTES:
        raise FileError("Файл больше 15 МБ")
    kind, text, image = "text", "", None
    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        for ti, t in enumerate(doc.tables):
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            if rows:
                text += f"\n\nТаблица {ti + 1}:\n" + "\n".join(" | ".join(r) for r in rows[:60])
    elif ext == ".doc":
        # старый двоичный Word: вытаскиваем читаемый текст (UTF-16LE и cp1251), без внешних утилит
        chunks = []
        for enc in ("utf-16le", "cp1251"):
            s = data.decode(enc, errors="ignore")
            chunks += [w for w in __import__("re").findall(r"[\wА-Яа-яЁёӘәІіҢңҒғҮүҰұҚқӨөҺһ.,:;%()\-–—\s]{12,}", s)]
        text = "\n".join(dict.fromkeys(c.strip() for c in chunks if c.strip()))
        if len(text) < 40:
            raise FileError("Не удалось прочитать .doc — сохраните документ как .docx или .pdf")
    elif ext == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(data))
        text = "\n".join(f"[стр. {i + 1}] " + (p.extract_text() or "") for i, p in enumerate(r.pages[:60]))
        if not text.strip():
            raise FileError("В PDF нет текстового слоя (скан). Пришлите изображение страницы — агент прочитает его")
    elif ext in (".xlsx", ".xls"):
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
        kind = "table"
        text = "\n\n".join(_table_text(df, f"{name} / {sh}") for sh, df in list(sheets.items())[:8])
    elif ext == ".csv":
        raw = data.decode("utf-8-sig", errors="ignore")
        sep = ";" if raw[:2000].count(";") > raw[:2000].count(",") else ","
        kind = "table"
        text = _table_text(pd.read_csv(io.StringIO(raw), sep=sep), name)
    elif ext == ".txt":
        text = data.decode("utf-8", errors="ignore")
    else:
        kind = "image"
        mime = "image/png" if ext == ".png" else "image/jpeg"
        image = f"data:{mime};base64," + base64.b64encode(data).decode()
        text = f"Изображение {name} ({len(data) // 1024} КБ)."
    return {"kind": kind, "text": text[:MAX_CHARS], "truncated": len(text) > MAX_CHARS, "image": image}


def save(session: str, name: str, data: bytes) -> dict:
    _gc()
    info = extract(name, data)
    fid = uuid.uuid4().hex[:12]
    _FILES[fid] = {"id": fid, "session": session, "name": name, "size": len(data), "ts": time.time(), **info}
    return {"id": fid, "name": name, "kind": info["kind"], "size": len(data), "chars": len(info["text"]),
            "truncated": info["truncated"], "preview": info["text"][:300]}


def get(session: str, ids: list[str]) -> list[dict]:
    return [_FILES[i] for i in ids if i in _FILES and _FILES[i]["session"] == session]


def _gc():
    now = time.time()
    for k in [k for k, v in _FILES.items() if now - v["ts"] > TTL_S]:
        _FILES.pop(k, None)
