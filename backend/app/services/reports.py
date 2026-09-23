"""Отчёты по ответу агента: Word (.docx) и Excel (.xlsx) с графиками и таблицами на реальных данных."""
import io
import re
from datetime import datetime, timezone

import pandas as pd

from app import db
from app.core import config
from app.services.alerts import build_alerts
from app.services.export import overall_metrics

LOCAL = pd.Timedelta(hours=config.SOURCE_UTC_OFFSET)
BLUE, AMBER, GREY = "#1f6feb", "#d29922", "#8b949e"


def _rated(turbine: str) -> float:
    return sum(t.rated_power_mw for t in config.TURBINES) if turbine == "STATION" else \
        next(t.rated_power_mw for t in config.TURBINES if t.id == turbine)


def _frame(issue_date: str, turbine: str) -> pd.DataFrame:
    store = db.get_store()
    run = store.latest_run(issue_date)
    if not run:
        return pd.DataFrame()
    fc = store.forecasts(run["id"], turbine).sort_values("target_time")
    rated = _rated(turbine)
    t = pd.to_datetime(fc["target_time"], utc=True) + LOCAL
    out = pd.DataFrame({
        "Время (Алматы)": t.dt.strftime("%d.%m %H:%M"),
        "Горизонт, ч": fc["lead_hours"].astype(int),
        "Прогноз, МВт": (fc["p_hat"].astype(float) * rated).round(3),
        "Факт, МВт": (fc["actual"].astype(float) * rated).round(3),
        "Ветер v_eq, м/с": fc["v_eq"].astype(float).round(2),
    })
    out["Отклонение, МВт"] = (out["Прогноз, МВт"] - out["Факт, МВт"]).round(3)
    return out.reset_index(drop=True)


def _name(turbine: str) -> str:
    return {"STATION": "ВЭС «Нурлы» (станция)", "T1": "Нурлы — турбина 1", "T2": "Нурлы — турбина 2"}[turbine]


def _chart_png(df: pd.DataFrame, title: str) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(9, 3.6), dpi=160)
    x = range(len(df))
    ax.fill_between(x, df["Прогноз, МВт"], color=BLUE, alpha=0.12)
    ax.plot(x, df["Прогноз, МВт"], color=BLUE, lw=2, ls="--", label="Прогноз")
    if df["Факт, МВт"].notna().any():
        ax.plot(x, df["Факт, МВт"], color=AMBER, lw=2, label="Факт")
    ax2 = ax.twinx()
    ax2.plot(x, df["Ветер v_eq, м/с"], color=GREY, lw=1, alpha=0.8, label="Ветер, м/с")
    ax2.set_ylabel("м/с", color=GREY)
    step = max(1, len(df) // 8)
    ax.set_xticks(list(x)[::step]); ax.set_xticklabels(df["Время (Алматы)"][::step], rotation=0, fontsize=8)
    ax.set_ylabel("МВт"); ax.set_title(title, fontsize=11, loc="left")
    ax.grid(alpha=0.25, ls=":"); ax.spines[["top"]].set_visible(False)
    h1, l1 = ax.get_legend_handles_labels(); h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, loc="upper right", fontsize=8, frameon=False)
    fig.tight_layout()
    buf = io.BytesIO(); fig.savefig(buf, format="png"); plt.close(fig)
    return buf.getvalue()


def _clean(text: str) -> list[str]:
    return [re.sub(r"\*\*(.+?)\*\*", r"\1", ln).strip() for ln in (text or "").splitlines()]


def build_docx(question: str, answer: str, issue_date: str, turbine: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Cm

    df = _frame(issue_date, turbine)
    rated = _rated(turbine)
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)

    h = doc.add_heading("Прогноз выработки ВЭС — ответ инженер-агента", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta = doc.add_paragraph()
    meta.add_run(f"{_name(turbine)} · выпуск {pd.Timestamp(issue_date):%d.%m.%Y} 05:00 (Алматы) · горизонт 48 ч · "
                 f"Pном {rated:g} МВт\nСформировано {datetime.now(timezone.utc) + LOCAL:%d.%m.%Y %H:%M} (Алматы), Infinity AI").font.color.rgb = RGBColor(0x59, 0x63, 0x6E)

    doc.add_heading("Вопрос", level=1)
    doc.add_paragraph(question or "—")
    doc.add_heading("Ответ агента", level=1)
    for ln in _clean(answer):
        if not ln:
            continue
        if ln.startswith(("- ", "• ", "* ")):
            doc.add_paragraph(ln[2:], style="List Bullet")
        elif re.match(r"^\d+[.)] ", ln):
            doc.add_paragraph(re.sub(r"^\d+[.)] ", "", ln), style="List Number")
        else:
            doc.add_paragraph(ln)

    if not df.empty:
        doc.add_heading("Ключевые показатели", level=1)
        fact = df["Факт, МВт"].dropna()
        kp = [("Энергия за 48 ч (прогноз)", f"{df['Прогноз, МВт'].sum():.1f} МВт·ч"),
              ("Средняя мощность", f"{df['Прогноз, МВт'].mean():.2f} МВт ({df['Прогноз, МВт'].mean() / rated:.0%} Pном)"),
              ("Пик прогноза", f"{df['Прогноз, МВт'].max():.2f} МВт в {df.loc[df['Прогноз, МВт'].idxmax(), 'Время (Алматы)']}"),
              ("Факт за горизонт", f"{fact.sum():.1f} МВт·ч ({len(fact)} ч)" if len(fact) else "факт ещё не наступил"),
              ("Средняя абсолютная ошибка", f"{df['Отклонение, МВт'].abs().mean():.2f} МВт" if len(fact) else "—")]
        t = doc.add_table(rows=0, cols=2); t.style = "Light Grid Accent 1"
        for k, v in kp:
            r = t.add_row().cells; r[0].text = k; r[1].text = v

        doc.add_heading("График: прогноз, факт и ветер", level=1)
        doc.add_picture(io.BytesIO(_chart_png(df, f"{_name(turbine)} · выпуск {pd.Timestamp(issue_date):%d.%m}")), width=Cm(17))

        alerts = build_alerts(issue_date, turbine)
        if alerts:
            doc.add_heading("Уведомления агента", level=1)
            for a in alerts:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run({"critical": "Критично. ", "warn": "Внимание. ", "info": ""}[a["level"]]).bold = True
                p.add_run(a["text"])

        doc.add_heading("Почасовой прогноз", level=1)
        tbl = doc.add_table(rows=1, cols=len(df.columns)); tbl.style = "Light List Accent 1"
        for i, c in enumerate(df.columns):
            tbl.rows[0].cells[i].text = c
        for _, row in df.iterrows():
            cells = tbl.add_row().cells
            for i, c in enumerate(df.columns):
                v = row[c]
                cells[i].text = "—" if pd.isna(v) else (f"{v:.2f}" if isinstance(v, float) else str(v))
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)

    doc.add_heading("Методика", level=2)
    doc.add_paragraph("Прогноз строится только из архивных прогнозов погоды, опубликованных до момента выпуска "
                      "(утечка будущего исключена). Модель: кривая мощности по прогнозному ветру 100 м с поправкой на "
                      "плотность воздуха и бустинг остатков. Факт показан только для оценки.")
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


def build_xlsx(question: str, answer: str, issue_date: str, turbine: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    df = _frame(issue_date, turbine)
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F6FEB"); head_font = Font(bold=True, color="FFFFFF")
    thin = Border(bottom=Side(style="thin", color="D0D7DE"))

    ws = wb.active; ws.title = "Ответ агента"
    ws["A1"] = "Прогноз выработки ВЭС — ответ инженер-агента"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{_name(turbine)} · выпуск {pd.Timestamp(issue_date):%d.%m.%Y} · горизонт 48 ч · Pном {_rated(turbine):g} МВт"
    ws["A2"].font = Font(color="59636E")
    ws["A4"] = "Вопрос"; ws["A4"].font = Font(bold=True); ws["A5"] = question
    ws["A7"] = "Ответ"; ws["A7"].font = Font(bold=True)
    for i, ln in enumerate([x for x in _clean(answer) if x], start=8):
        ws.cell(row=i, column=1, value=ln).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120

    if not df.empty:
        wf = wb.create_sheet("Прогноз по часам")
        for j, c in enumerate(df.columns, start=1):
            cell = wf.cell(row=1, column=j, value=c); cell.fill = head_fill; cell.font = head_font
            wf.column_dimensions[get_column_letter(j)].width = 18
        for i, row in enumerate(df.itertuples(index=False), start=2):
            for j, v in enumerate(row, start=1):
                c = wf.cell(row=i, column=j, value=None if pd.isna(v) else v); c.border = thin
                if isinstance(v, float):
                    c.number_format = "0.00"
        wf.freeze_panes = "A2"
        n = len(df) + 1
        ch = LineChart(); ch.title = "Прогноз и факт, МВт"; ch.y_axis.title = "МВт"; ch.height, ch.width = 9, 26
        ch.add_data(Reference(wf, min_col=3, max_col=4, min_row=1, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(wf, min_col=1, min_row=2, max_row=n))
        ch.series[0].graphicalProperties.line.solidFill = "1F6FEB"; ch.series[0].graphicalProperties.line.dashStyle = "dash"
        if len(ch.series) > 1:
            ch.series[1].graphicalProperties.line.solidFill = "D29922"
        wf.add_chart(ch, "H2")
        cw = LineChart(); cw.title = "Ветер v_eq, м/с"; cw.height, cw.width = 7, 26
        cw.add_data(Reference(wf, min_col=5, min_row=1, max_row=n), titles_from_data=True)
        cw.set_categories(Reference(wf, min_col=1, min_row=2, max_row=n))
        cw.series[0].graphicalProperties.line.solidFill = "8B949E"
        wf.add_chart(cw, "H22")

        wa = wb.create_sheet("Уведомления")
        for j, c in enumerate(["Уровень", "Тип", "Начало (UTC)", "Конец (UTC)", "Изменение, МВт", "Сообщение"], start=1):
            cell = wa.cell(row=1, column=j, value=c); cell.fill = head_fill; cell.font = head_font
        for i, a in enumerate(build_alerts(issue_date, turbine), start=2):
            for j, v in enumerate([a["level"], a["kind"], a["start"], a["end"], a.get("delta_mw"), a["text"]], start=1):
                wa.cell(row=i, column=j, value=v).alignment = Alignment(wrap_text=j == 6, vertical="top")
        for col, w in zip("ABCDEF", (10, 12, 22, 22, 14, 90)):
            wa.column_dimensions[col].width = w

    m = overall_metrics(db.get_store().latest_forecasts())
    daily = [d for d in m.get("daily", []) if d["turbine"] == turbine]
    if daily:
        wm = wb.create_sheet("Точность (январь)")
        cols = ["Дата выпуска", "MAE базы", "MAE модели 1–24 ч", "MAE модели 25–48 ч"]
        for j, c in enumerate(cols, start=1):
            cell = wm.cell(row=1, column=j, value=c); cell.fill = head_fill; cell.font = head_font
            wm.column_dimensions[get_column_letter(j)].width = 20
        for i, d in enumerate(daily, start=2):
            for j, v in enumerate([d["issue_date"], d["mae_base"], d["mae_24"], d["mae_48"]], start=1):
                wm.cell(row=i, column=j, value=v).number_format = "0.000"
        bc = BarChart(); bc.title = "Ошибка по дням: база против модели"; bc.height, bc.width = 9, 26
        bc.add_data(Reference(wm, min_col=2, max_col=4, min_row=1, max_row=len(daily) + 1), titles_from_data=True)
        bc.set_categories(Reference(wm, min_col=1, min_row=2, max_row=len(daily) + 1))
        for s, col in zip(bc.series, ("C8D1DA", "1F6FEB", "D29922")):
            s.graphicalProperties.solidFill = col
        wm.add_chart(bc, "F2")
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
